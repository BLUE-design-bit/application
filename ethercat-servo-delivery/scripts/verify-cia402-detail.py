"""Read-only checks: preserved initial workbook, complete table/type inventory, export fidelity."""
from pathlib import Path
from copy import copy
import json,re,collections
from openpyxl import load_workbook
root=Path(__file__).resolve().parents[1]
out=root/'outputs/01a08f03-8e5c-79e3-bb07-33816e14789f'
before=load_workbook(out/'DFS10A_CiA402_V3_模块覆盖初审.xlsx')
after=load_workbook(out/'DFS10A_CiA402_V3_数据类型逐项核对_v2.xlsx')
data=json.loads((root/'.cache/cia402-detail/detail-data.json').read_text(encoding='utf-8'))
assert after.sheetnames[:3]==before.sheetnames
for old in before:
    new=after[old.title]
    assert (new.max_row,new.max_column)==(old.max_row,old.max_column)
    for row in old:
        for c in row:
            x=new[c.coordinate]
            assert c.value==x.value,(old.title,c.coordinate,'value')
            for attr in ['font','fill','alignment','border','protection']:
                assert copy(getattr(c,attr))==copy(getattr(x,attr)),(old.title,c.coordinate,attr)
            assert c.number_format==x.number_format
    assert old.freeze_panes==new.freeze_panes
    assert list(old.tables)==list(new.tables)
    for col,dim in old.column_dimensions.items():assert dim.width==new.column_dimensions[col].width
    for row,dim in old.row_dimensions.items():assert dim.height==new.row_dimensions[row].height
    assert len(old.conditional_formatting)==len(new.conditional_formatting)
names={'types':'数据类型逐项'}
assert len(after.sheetnames)==4
for key,name in names.items():
    sh=after[name];rows=data[key]
    assert sh.max_row==len(rows)+6
    for i,expected in enumerate(rows,7):
        for j,value in enumerate(expected,1):
            actual=sh.cell(i,j).value
            assert actual==(None if value=='' else value),(name,i,j,actual,value)
    assert sh.freeze_panes=='C7' and len(sh.tables)==1
    assert not any(c.data_type=='e' for row in sh for c in row)
    assert max(d.height or 0 for d in sh.row_dimensions.values())<=409
    if key!='manifest':
        assert len(sh.data_validations.dataValidation)==1
        checkcol=5 if key=='corrigendum' else 10
        assert all(sh.cell(r,checkcol).value=='未核对' for r in range(7,sh.max_row+1))
assert len(data['types'])==24
assert [sum(r[0].startswith('T'+str(i).zfill(2)+'-') for r in data['types']) for i in range(1,5)]==[11,3,7,3]
print('PASS: 3 original sheets preserved including styles; chapter 5 all 24 checks; export fidelity; dropdowns; no error cells.')
