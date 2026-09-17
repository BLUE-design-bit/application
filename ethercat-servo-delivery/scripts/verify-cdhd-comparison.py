"""Read-only checks for the CDHD declaration comparison and preserved audit."""
from pathlib import Path
from copy import copy
import json,re
from lxml import etree
from openpyxl import load_workbook

root=Path(__file__).resolve().parents[1]
folder=root/'outputs/01a08f03-8e5c-79e3-bb07-33816e14789f'
data=json.loads((root/'.cache/cia402-detail/cdhd-comparison.json').read_text(encoding='utf-8'))
before=load_workbook(folder/'DFS10A_CiA402_V3_全章节逐项核对_v3.xlsx')
output=folder/'DFS10A_高创CDHD2_CiA402逐项对照_v4.xlsx'
after=load_workbook(output)
cached=load_workbook(output,data_only=True)
assert after.sheetnames[:9]==before.sheetnames and len(after.sheetnames)==12
count=0
for old in before:
    new=after[old.title];cfg=data['changed'].get(old.title)
    assert old.max_row==new.max_row
    assert new.max_column==old.max_column+(2 if cfg else 0)
    assert old.freeze_panes==new.freeze_panes
    assert list(old.tables)==list(new.tables)
    for row in old:
        for c in row:
            target=c.column+(2 if cfg and c.column>cfg['position']+1 else 0)
            n=new.cell(c.row,target)
            expected=c.value
            if cfg and c.row==6:expected=cfg['headers'][target-1]
            assert n.value==expected,(old.title,c.coordinate,n.coordinate,n.value,expected)
            for field in ['font','fill','alignment','border','protection']:
                # "general" versus omitted horizontal alignment has the same Excel meaning.
                a,b=copy(getattr(c,field)),copy(getattr(n,field))
                if field=='alignment':
                    if a.horizontal is None:a.horizontal='general'
                    if b.horizontal is None:b.horizontal='general'
                    if a.wrapText is None:a.wrapText=False
                    if b.wrapText is None:b.wrapText=False
                assert a==b,(old.title,c.coordinate,n.coordinate,field,a,b)
            if not cfg:assert c.number_format==n.number_format
            count+=1
    if cfg:
        for r,expected in enumerate(cfg['rows'],7):
            actual=[new.cell(r,c).value for c in range(1,len(expected)+1)]
            assert actual==[None if v=='' else v for v in expected],(old.title,r,'export values')
        validations=new.data_validations.dataValidation
        if cfg['oldCols']==11:
            assert len(validations)==1,(old.title,[(str(v.sqref),v.formula1) for v in validations])
            assert str(validations[0].sqref)==f'L7:L{cfg["last"]}'
            assert validations[0].formula1=='"未核对,已核对,有异议,待实测"'
        else:assert not validations
    else:
        assert old.data_validations==new.data_validations
        for col,dim in old.column_dimensions.items():assert dim.width==new.column_dimensions[col].width
        for r,dim in old.row_dimensions.items():assert dim.height==new.row_dimensions[r].height
    assert not any(c.data_type=='e' for row in new for c in row)

for name,key in [('参考对象对照','inventory'),('高创声明差异','issues')]:
    s=after[name]
    for r,expected in enumerate(data[key],7):
        actual=[s.cell(r,c).value for c in range(1,len(expected)+1)]
        assert actual==expected,(name,r,actual,expected)
    assert s.max_row==len(data[key])+6
assert len(data['inventory'])==103 and len(data['issues'])==9
s=cached['高创参考覆盖']
for cell,value in [('B7',72),('B8',23),('B9',49),('B10',23/72),('B11',1)]:
    assert s[cell].value==value,(cell,s[cell].value,value)
    assert after[s.title][cell].data_type=='f'
assert s['B26'].value=='0x1DD'
assert not any(c.data_type=='e' for sheet in cached for row in sheet for c in row)

# Re-extract keys from XML and compare against actual C object registrations,
# independent of the prepared coverage flags and parser's decoded defaults.
xml=etree.parse(data['source']['path'],etree.XMLParser(resolve_entities=False,no_network=True))
objs={int(o.findtext('Index').replace('#x','0x'),0):o for o in xml.findall('.//Objects/Object')}
assert len(objs)==540
inventory=after['参考对象对照']
std={int(inventory.cell(r,2).value,16) for r in range(7,110)}
code=root/'dfs-10a/df-sdk/modules/ethercat'
regs=set()
for file in ['cia402appl.h','coeappl.c']:
    source=(code/file).read_text(encoding='utf-8',errors='replace')
    regs.update(int(x,16) for x in re.findall(r'\{\s*(?:NULL|0)\s*,\s*(?:NULL|0)\s*,\s*0x([0-9a-fA-F]{4})\s*,',source))
assert len(regs)>20,'Registration pattern must match actual source'
assert len(std&objs.keys())==72
assert len(std&objs.keys()&regs)==23
for r in range(7,110):
    index=int(inventory.cell(r,2).value,16)
    assert inventory.cell(r,5).value==int(index in regs)
    assert inventory.cell(r,6).value==int(index in objs)
mask=int(objs[0x6502].findtext('Info/DefaultValue').replace('#x','0x'),0)
assert mask==0x1DD and mask&(1<<4) and not mask&(1<<5) and not mask&(1<<9)
assert all(int(objs[x].findtext('Info/DefaultValue').replace('#x','0x'),0)==65535 for x in [0x605B,0x605C,0x605E])
types={x.findtext('Name'):x for x in xml.findall('.//DataTypes/DataType')}
assert types['DT608F'].find("SubItem[SubIdx='2']/Flags/Access").text=='ro'
assert types['DT60FE'].find("SubItem[SubIdx='1']/Type").text=='UDINT'
assert types['DT607DARR'].findtext('BaseType')=='DINT'
assert types['DT607DARR'].findtext('ArrayInfo/Elements')=='2'
print(f'PASS: {count} original cells preserved with equivalent styles; all added judgments exported exactly; manual checks moved to L only; 103 objects/9 differences; XML/source independently reconcile 23/72 and 49 missing; cached formulas and error scan passed.')
