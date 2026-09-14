"""Read-only verification of the exported blank functional checklist."""
import json
import re
from pathlib import Path
from openpyxl import load_workbook

root = Path(__file__).resolve().parent.parent
file = root / 'outputs/01a08f03-8e5c-79e3-bb07-33816e14789f/EtherCAT_功能与路径实测_v2.xlsx'
scope = (root / 'functional-scope-review.md').read_text(encoding='utf-8')
source = []
paths = []
group = ''
for line in scope.splitlines():
    h = re.match(r'^## \d+\. (.+)：已确认$', line)
    if h:
        group = h[1]
    m = re.match(r'^\| ([MDSUHNPAECTR]\d{2}) \| (.+) \| (.+) \|$', line)
    if m:
        (paths if m[1].startswith('E') else source).append((group, *m.groups()))
assert len(source) == 106 and len(paths) == 12
assert not any(x[1] == 'M08' for x in source)
wb = load_workbook(file, data_only=False)
cached = load_workbook(file, data_only=True)
assert wb.sheetnames == ['品牌勾选', '路径与实测记录', '设备与测试条件']
main, log, config = [wb[n] for n in wb.sheetnames]
mc, lc, cc = [cached[n] for n in cached.sheetnames]
platforms = ['汇川小型', '汇川中大型', '欧姆龙']
brands = ['CDHD', '台达', '繁易', 'Elmo', 'KEBA（可选）']
assert [main.cell(9,c).value for c in range(6,11)] == brands
for p, plc in enumerate(platforms):
    for i, (section, ident, name, coverage) in enumerate(source):
        r = 10 + p * 106 + i
        assert [main.cell(r,c).value for c in range(1,6)] == [plc, section, ident, name, coverage]
        assert all(main.cell(r,c).value == '☐ 未测试' for c in range(6,11))
assert main.tables['FunctionMatrix'].ref == 'A9:J327'
assert log.tables['PathRecords'].ref == 'A9:U1600'
assert config.tables['Equipment'].ref == 'A9:L24'
assert config.tables['Conditions'].ref == 'A28:H43'
assert main.freeze_panes == 'D10'
assert log.freeze_panes == 'C10'
assert config.freeze_panes == 'B10'
for sh in wb:
    assert sh.sheet_view.showGridLines is False
    assert sh.data_validations.count > 0
    assert len(sh.conditional_formatting) > 0

for n, (p, b) in enumerate((p,b) for p in range(3) for b in range(5)):
    combo = f'P{p+1}-B{b+1}-01'
    r = 10+n
    assert [config.cell(r,c).value for c in range(1,4)] == [combo, platforms[p], brands[b]]
    assert all(config.cell(r,c).value is None for c in range(4,12))
    assert cc.cell(r,12).value == '待补型号/版本'
    for f, (_, ident, _, _) in enumerate(source):
        row = 10+n*106+f
        assert log.cell(row,1).value == f'REC-{n*106+f+1:05d}'
        assert log.cell(row,2).value == combo
        assert lc.cell(row,3).value == platforms[p]
        assert lc.cell(row,4).value == brands[b]
        assert log.cell(row,5).value == ident
        for c in [6,7,9,*range(11,18),19,20]:
            assert log.cell(row,c).value is None, (row,c)
        assert log.cell(row,8).value == '待判定'
        assert log.cell(row,10).value == '未测试'
        assert log.cell(row,18).value == '待展开'
        assert lc.cell(row,21).value == '待展开路径'
        assert log.cell(row,15).number_format == 'yyyy-mm-dd'
        assert log.cell(row,20).number_format == 'yyyy-mm-dd'
        for c in [3,4,21]:
            assert log.cell(row,c).data_type == 'f'
            assert f'B{row}' in log.cell(row,c).value

for c in range(6,11):
    assert mc.cell(7,c).value == mc.cell(8,c).value == 0
assert lc.cell(1600,21).value in (None,'')
for c in [1,2,*range(5,21)]:
    assert log.cell(1600,c).value is None
assert any('INDIRECT("Equipment[组合编号]")' in (v.formula1 or '') for v in log.data_validations.dataValidation)
assert any('部分实测' in (v.formula1 or '') for v in main.data_validations.dataValidation)
for i, (_, ident, name, coverage) in enumerate(paths):
    assert config.cell(49+i,1).value == ident
    assert config.cell(49+i,2).value == name
    assert config.cell(49+i,4).value == coverage
for sh in cached:
    for row in sh:
        for cell in row:
            assert cell.data_type != 'e', (sh.title, cell.coordinate, cell.value)
            if isinstance(cell.value,str):
                assert '验证用' not in cell.value, (sh.title,cell.coordinate)
assert not wb._external_links
print(json.dumps({'verified': str(file), 'functions': 106, 'paths':12, 'platforms':platforms,
                  'matrix_rows':318, 'untested_brand_cells':1590, 'blank_planned_records':1590,
                  'device_placeholders':15, 'sheets':wb.sheetnames,
                  'checks':'source fidelity, blank evidence, cached formulas, extension row, table refs, validations, freeze panes, dates'},ensure_ascii=False))
