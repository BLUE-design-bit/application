"""Read-only verification of the exported initial CiA402 audit workbook."""
from pathlib import Path
from openpyxl import load_workbook

root = Path(__file__).resolve().parents[1]
book = root / 'outputs/01a08f03-8e5c-79e3-bb07-33816e14789f/DFS10A_CiA402_V3_模块覆盖初审.xlsx'
wb = load_workbook(book, data_only=False)
assert wb.sheetnames == ['目录模块总览', '重点差异与缺口', '审查依据']
main, detail, source = wb.worksheets
assert [main.cell(r, 1).value for r in range(7, 27)] == list(map(str, range(1, 21)))
assert main['A27'].value == '勘误1'
assert detail.max_row == 18 and source.max_row == 20
ids = {source.cell(r, 1).value for r in range(7, 21)}
for sheet in [main, detail]:
    for row in sheet.iter_rows(min_row=7):
        assert set(row[7].value.split('/')) <= ids
for sheet in wb:
    assert sheet.freeze_panes == 'C7'
    assert len(sheet.tables) == 1
    assert not any(c.data_type == 'e' for row in sheet for c in row)
    assert all(c.value is not None for row in sheet.iter_rows(min_row=6) for c in row)
assert sum(detail.cell(r, 3).value == '明确差异' for r in range(7, 19)) == 2
print('PASS: 20 chapters + corrigendum; 12 findings; 14 evidence records; complete references, tables and panes.')
