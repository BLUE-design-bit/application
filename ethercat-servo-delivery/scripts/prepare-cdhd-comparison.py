"""Compare the customer CDHD2 ESI declarations to the accepted CiA402 audit.

XML declarations are evidence of declaration only, not firmware conformance.
"""
from pathlib import Path
from lxml import etree as ET
from openpyxl import load_workbook
import json,re,hashlib,collections

root=Path(__file__).resolve().parents[1]
cache=root/'.cache/cia402-detail'
xml=Path('C:/Users/Admin/Downloads/Servotronix_CDHD2_PC0_RV0_12C18003_2_38_6C1_20230605.xml')
doc=ET.parse(str(xml),ET.XMLParser(resolve_entities=False,no_network=True))
devices=doc.findall('.//Device');assert len(devices)==1
d=devices[0];dtypes={x.findtext('Name'):x for x in d.findall('.//DataTypes/DataType')}
def integer(s):
    if s is None:return None
    s=s.strip().replace(' ','')
    try:
        if s.lower().startswith('#x'):return int(s[2:],16)
        if s.lower().startswith('0x'):return int(s,16)
        if s.lower().endswith('h'):return int(s[:-1],16)
        return int(s)
    except ValueError:return None
aliases={'USINT':'Unsigned8','UINT':'Unsigned16','UDINT':'Unsigned32','SINT':'Integer8','INT':'Integer16','DINT':'Integer32','BYTE':'Unsigned8','WORD':'Unsigned16','DWORD':'Unsigned32'}
def typename(t):return aliases.get(t,'Visible string' if t.startswith('STRING') else t)
def scalar(e):
    typ=e.findtext('Type','');bits=integer(e.findtext('BitSize'))
    return dict(type=typename(typ),rawType=typ,bits=bits,access=e.findtext('Flags/Access'),pdo=e.findtext('Flags/PdoMapping'),min=e.findtext('Info/MinValue'),max=e.findtext('Info/MaxValue'),line=e.sourceline,name=e.findtext('Name','').strip())
def decode(value,typ,bits):
    n=integer(value)
    if n is None:return value
    if typ.startswith('Integer') and bits and n>=2**(bits-1):n-=2**bits
    return n
objects={}
for o in d.findall('.//Objects/Object'):
    idx=f'{integer(o.findtext("Index")):04X}';ob=scalar(o);ob['index']=idx;ob['fields']={};ob['objectLine']=o.sourceline
    dt=dtypes.get(o.findtext('Type'));rawfields=[]
    if dt is not None:
        for s in dt.findall('SubItem'):
            si=integer(s.findtext('SubIdx'))
            if si is not None:rawfields.append((si,scalar(s)))
            else:
                at=dtypes.get(s.findtext('Type'));arr=at.find('ArrayInfo') if at is not None else None
                assert arr is not None,(idx,s.findtext('Type'))
                lower=integer(arr.findtext('LBound'));count=integer(arr.findtext('Elements'))
                for k in range(lower,lower+count):
                    fld=scalar(s);fld.update(type=typename(at.findtext('BaseType')),rawType=at.findtext('BaseType'),bits=integer(at.findtext('BitSize'))//count,arrayLine=at.sourceline,name=f'Sub-Index {k}')
                    rawfields.append((k,fld))
    info=o.findall('Info/SubItem')
    if rawfields:
        for order,(si,fld) in enumerate(rawfields):
            matches=[it for it in info if it.findtext('Name','').strip()==fld['name']]
            it=matches[0] if len(matches)==1 else (info[order] if order<len(info) else None)
            default=it.findtext('Info/DefaultValue') if it is not None else None
            fld['defaultRaw']=default;fld['default']=decode(default,fld['type'],fld['bits']);fld['defaultLine']=it.sourceline if it is not None else None
            ob['fields'][str(si)]=fld
        ob['shape']='带子项结构'
    else:
        fld=dict(ob);fld.pop('fields');raw=o.findtext('Info/DefaultValue')
        fld['defaultRaw']=raw;fld['default']=decode(raw,fld['type'],fld['bits']);fld['defaultLine']=o.sourceline
        ob['fields']['0']=fld;ob['shape']='标量'
    assert idx not in objects
    objects[idx]=ob

mask=objects['6502']['fields']['0']['default'];assert mask==0x1dd
validbits={0:'PP',1:'vl',2:'PV',3:'PT',5:'HM',6:'IP',7:'CSP',8:'CSV',9:'CST'}
def evidence(idx,sub=None):
    if idx not in objects:return f'XML Dictionary/Objects：未列出0x{idx}；不据此断言设备不支持'
    ob=objects[idx];out=f'XML Object 0x{idx}：L{ob["objectLine"]}'
    if sub is not None and str(sub) in ob['fields']:
        f=ob['fields'][str(sub)];out+=f'；子项{sub:02X}类型/权限：L{f["line"]}'
        if f.get('arrayLine'):out+=f'；数组展开：L{f["arrayLine"]}'
        if f.get('defaultLine'):out+=f'；默认值：L{f["defaultLine"]}'
    return out
def brief(idx,sub=None):
    ob=objects[idx]
    if sub is None and ob['shape']!='标量':return f'{ob["rawType"]}，列出子项'+','.join(f'{int(s):02X}' for s in ob['fields'])
    f=ob['fields'].get(str(sub or 0));
    if not f:return '未列出该子项'
    return f'{f["type"]}/{f["bits"]}bit，{f["access"] or "未列访问权限"}，默认{f["default"] if f["default"] is not None else "未列"}'
full=json.loads((cache/'full-detail-data.json').read_text(encoding='utf-8'))
ours=full['code'];standard={}
for r in full['attributes']:
    if r[3].endswith('对象索引'):standard[r[3].split()[0]]=dict(section=r[1],name='',category=r[5])
    if r[3].endswith('对象名称'):standard[r[3].split()[0]]['name']=r[4].split('（')[0]
for idx,label in [('1001','Error register'),('1018','Identity')]:standard[idx]=dict(section='6.2',name=label,category='正文6.2规定必须提供；详细定义引用CiA301')
assert len(standard)==103
section_obj={v['section']:idx for idx,v in standard.items() if idx not in ['1000','1001','1018']}

def declared(idx,sub=None):
    if idx not in objects:return 'XML未列出；不能认定设备缺失',evidence(idx)
    if sub is not None and str(sub) not in objects[idx]['fields']:return f'XML未列子项{sub:02X}；不等于实机缺失',evidence(idx)
    return 'XML已声明：'+brief(idx,sub)+'；行为须实测',evidence(idx,sub)

def attr_judgment(r):
    m=re.match(r'([0-9A-F]{4})(?::([0-9A-F]{2}))? (.*)',r[3]);assert m,r
    idx,sub,key=m.groups();si=int(sub,16) if sub else None
    if idx not in objects:return declared(idx,si)
    ob=objects[idx];req=r[4].split('（')[0];ev=evidence(idx,si)
    if si is not None and str(si) not in ob['fields']:return declared(idx,si)
    f=ob['fields'].get(str(si or 0));msg=brief(idx,si)
    if key in ['对象索引','对象名称','子索引','对象实现要求','字段实现要求']:return 'XML已声明：'+msg+'；不代表运行通过',ev
    if key=='对象结构':
        if req=='Variable':return ('声明结构一致：标量' if ob['shape']=='标量' else '声明结构有差异：XML为带子项结构'),ev
        if req in ['Array','Record']:return ('已声明子项；ObjectCode未明示，不能仅据DT名称判Array/Record' if ob['shape']!='标量' else '声明结构有差异：XML为标量'),ev
        return 'XML：'+msg+'；标准引用CiA301，未扩展验证',ev
    if key=='数据类型':
        if ob['shape']!='标量':return '已列出结构字段；逐字段见数据类型/子项检查，不能按总BitSize判单个字段',ev
        return ('类型声明一致：' if req.lower().replace(' ','')==f['type'].lower().replace(' ','') else '类型声明需核对：')+msg,ev
    if key=='访问权限':
        if req=='c' and f['access']=='ro':return '声明只读；是否恒定不能由ro单独证明',ev
        return ('访问声明一致：' if req==f['access'] else '访问声明有差异：')+f'标准{req}；XML {f["access"]}',ev
    if key=='默认值':
        actual=f['default'];expected=integer(req)
        if idx=='6060':expected=0
        if actual is None:return 'XML未给出默认值',ev
        if expected is not None:
            note='；上电实值待读回'
            if idx in ['605B','605C','605E'] and actual==-1:note+='；-1属于允许的厂商范围，本行只比较默认值'
            if idx=='60C4' and si==0:note+='；客户表3只列00–06，表168却写07，标准自身一致性也需核实'
            return ('默认声明一致：' if actual==expected else '默认声明有差异：')+f'标准{expected}；XML {actual}（{f["defaultRaw"]}）'+note,ev
        return f'XML默认{actual}（{f["defaultRaw"]}）；标准为{req}，不作固定值冲突判定',ev
    if key=='允许范围':
        if req==f['type']:return '位宽/符号声明一致：'+msg+'；未验证边界行为',ev
        return f'XML类型{f["type"]}；Min={f["min"] or "未列"}，Max={f["max"] or "未列"}；不能从默认值推断允许范围',ev
    if key=='PDO映射要求':return f'XML PdoMapping={f["pdo"] or "未列"}；映射规范引用Part 3，此处不判完整符合',ev
    return declared(idx,si)

chapter_judgments={
'1':'本次只依据该型号/版本ESI；不是高创全系列或实机认证结论',
'2':'XML不是CiA301/402-3全文，引用条款和运行测试未验证',
'3':'术语与标准一致使用；不属于功能符合性结论',
'4':'ProfileNo=402；6502=0x1DD，bit4保留位置1，模式声明有疑点',
'5':'基本整数/字符串已使用；60C2/60C4字段已列；Time of day及vl记录未见声明',
'6':'1000/1001/1018已列；6402–6407、6503/6505未列，未列不等于实机不支持',
'7':'603F/1001已列；错误码触发、清除和EMCY行为无法由ESI证明',
'8':'6040/6041/6060/6061/6502已列；605A未列；默认值及6502声明存在与客户标准的差异',
'9':'608F/6091/6092/607E已列；6090未列；换算行为待实测',
'10':'6502 bit0=1声明PP；607A等对象已列；轨迹/握手待实测',
'11':'607C/6098/6099/609A已列，但6502 bit5=0；回零能力声明需与实机核对',
'12':'位置反馈、跟随误差、位置窗口对象已列；闭环和状态触发须实测',
'13':'6502 bit6=1声明IP；60C0/C1/C2/C4已列；60C1只列01–04，非全部254子项',
'14':'6502 bit2=1声明PV；速度目标、反馈、窗口对象已列；动态行为待实测',
'15':'6502 bit3=1声明PT；转矩目标、反馈、额定值、斜坡对象已列；单位换算须实测',
'16':'6502 bit1=0，vl模式未声明；6042–604C未列',
'17':'6502 bit7=1声明CSP；位置偏置60B0/速度偏置60B1/转矩偏置60B2已列；动态行为待实测',
'18':'6502 bit8=1声明CSV；60FF/606C已列；速度单位及同步行为待实测',
'19':'6502 bit9=0，未声明CST；6071/6077同时用于PT，存在这些对象不能证明CST支持',
'20':'60FD已列；60FE列00/01/02及输出掩码，子项结构比我方VAR:00完整；实际I/O待实测',
}
def generic(r):
    sec=r[1];idx=section_obj.get(sec,'')
    m=re.match(r'([0-9A-F]{4})(?::([0-9A-F]{2}))?\b',r[3])
    if m:idx=m[1]
    if r[0].startswith('F07-'):
        bit=int(r[0][-2:]);v=(mask>>bit)&1
        return f'XML默认bit{bit}={v}；'+('保留位应为0，声明有差异' if v and bit in [4,*range(10,16)] else '仅表示支持声明，实际能力待实测'),evidence('6502',0)
    if r[0].startswith(('F05-','F06-')):idx='6040' if r[0].startswith('F05') else '6041'
    if r[0].startswith('F71-'):idx='60FD'
    if r[0].startswith('F72-'):idx='60FE'
    if r[0].startswith('T053-'):
        v=integer(r[3].split(' / ')[-1])
        if v in range(1,11):return f'6502 bit{v-1}={(mask>>(v-1))&1}；'+('标准保留位被置1' if v==5 else '模式运行行为待实测'),evidence('6502',0)
    if sec=='6.2':
        m=re.search(r'\b(1000|1001|1018)\b',r[4]);idx=m[1] if m else idx
    if idx:return declared(idx)
    if sec.startswith('7'):return 'XML不能证明本条故障/错误码行为；603F及1001已列',evidence('603F')
    if r[6]=='勘误已删除' or r[6]=='删除条款':return '此条被客户附带勘误删除，不作为有效要求','客户PDF勘误；与品牌无关'
    ch=sec.split('.')[0]
    return chapter_judgments.get(ch,'XML只能确认声明；本条行为/图示无法据此判定'),evidence('6502',0) if ch in [str(i) for i in range(8,20)] else 'XML Profile/Dictionary；见高创参考覆盖中的基线'

def type_judgment(r):
    rid=r[0]
    if rid.startswith('T01-'):
        n=int(rid[-2:]);basic=['Unsigned8','Unsigned16','Unsigned32','Integer8','Integer16','Integer32','Visible string']
        if n<=7:
            target=basic[n-1];examples=[(i,s,f) for i,o in objects.items() for s,f in o['fields'].items() if f['type']==target]
            if examples:
                i,s,f=examples[0];return f'XML有实际对象使用{target}：{i}:{int(s):02X}；传输行为待实测',evidence(i,int(s))
            return f'XML未见对象使用{target}','XML Dictionary/DataTypes及Objects'
        if n==8:return 'XML未列Time of day；6406未列，不能断言固件不支持',evidence('6406')
        if n in [9,10]:return declared('60C2' if n==9 else '60C4')
        return 'XML未列6048/6049/604A；未声明vl模式',evidence('6502',0)
    obj='60C2' if rid.startswith('T02') else ('60C4' if rid.startswith('T03') else '6048')
    si=int(rid[-2:]);x,e=declared(obj,si)
    if obj in objects:
        expected=re.search(r';\s*(Unsigned\d+|Integer\d+)',r[4].replace('；',';'))
        actual=objects[obj]['fields'].get(str(si),{}).get('type')
        if expected and actual:x=('字段类型声明一致：' if expected[1]==actual else '字段类型声明有差异：')+brief(obj,si)+'；运行行为待实测'
    return x,e

book=load_workbook(root/'outputs/01a08f03-8e5c-79e3-bb07-33816e14789f/DFS10A_CiA402_V3_全章节逐项核对_v3.xlsx')
changed={};pos={'目录模块总览':4,'重点差异与缺口':2,'数据类型逐项':5,'对象属性逐项':6,'行为与位定义逐项':6,'正文条款逐项':6,'勘误逐项':6}
gap_objects={'D01':'605A','D02':'60FE','D03':'6077','D04':'6071','D05':'603F','D06':'6040','D07':'6041','D08':'606C','D09':'607D','D10':'60C2'}
for sheet,p in pos.items():
    sh=book[sheet];headers=[c.value for c in sh[6]];rows=[list(r) for r in sh.iter_rows(min_row=7,values_only=True)];new=[]
    for r in rows:
        if sheet=='目录模块总览':
            text=chapter_judgments.get(str(r[0]),'本列是XML声明判断；无法确认高创完全符合');ev=evidence('6502',0)
        elif sheet=='重点差异与缺口':
            ids=re.findall(r'\b[0-9A-F]{4}\b',r[3]);idx=ids[0] if ids else gap_objects.get(r[0]);text,ev=declared(idx) if idx else ('XML不能证明该项实际行为；需实机对照','见相关章节的逐项比较')
            if r[0]=='D01':text='605A未列；605B/C/E默认-1（厂商范围），与客户版默认0/1/2不同；动作无法由XML验证';ev='XML L11932、11944、11968；605A未列'
            if r[0]=='D11':text='声明PP/PV/PT/IP；vl未声明；HM对象存在但6502 bit5=0，需核实';ev=evidence('6502',0)
            if r[0]=='D12':text='608F/6091/6092/607E、60B8–60BD、60B0–60B2均已列；执行行为待实测';ev='XML对应Object索引；见参考对象对照'
        elif sheet=='数据类型逐项':text,ev=type_judgment(r)
        elif sheet=='对象属性逐项':text,ev=attr_judgment(r)
        else:text,ev=generic(r)
        new.append(r[:p+1]+[text,ev]+r[p+1:])
    headers[p]='我们的当前判断';headers[p+1:p+1]=['高创的判断（仅XML）','高创XML依据']
    if sheet in ['对象属性逐项','行为与位定义逐项','正文条款逐项','勘误逐项']:headers[5]='适用条件（我方当前模式）'
    formats=[]
    for column in range(p+2,sh.max_column+1):
        for row in range(1,sh.max_row+1):
            c=sh.cell(row,column)
            font=dict(name=c.font.name or 'Arial',size=c.font.sz or 10,bold=bool(c.font.b),italic=bool(c.font.i))
            if c.font.color and c.font.color.type=='rgb':font['color']='#'+c.font.color.rgb[-6:]
            fmt=dict(font=font,wrapText=bool(c.alignment.wrapText),verticalAlignment=c.alignment.vertical or 'bottom',horizontalAlignment=c.alignment.horizontal or 'general')
            if c.fill.patternType=='solid':fmt['fill']='#'+c.fill.fgColor.rgb[-6:]
            else:fmt['fill']=None
            if c.border.bottom and c.border.bottom.style:fmt['borders']={'bottom':{'style':c.border.bottom.style,'color':'#'+c.border.bottom.color.rgb[-6:]}}
            formats.append(dict(row=row,col=column+2,format=fmt,numberFormat=c.number_format))
    changed[sheet]=dict(position=p,headers=headers,rows=new,formats=formats,oldCols=sh.max_column,last=sh.max_row,widths=[sh.column_dimensions[c.column_letter].width for c in sh[6]],heights={str(i):v.height for i,v in sh.row_dimensions.items()},table=next(iter(sh.tables)),tableStyle=next(iter(sh.tables.values())).tableStyleInfo.name)

inventory=[]
for idx,v in sorted(standard.items(),key=lambda x:tuple(map(int,x[1]['section'].split('.')))):
    op=idx in ours;xp=idx in objects
    current='未在实际CoE字典注册' if not op else '对象已注册；细项差异/行为见原明细'
    if idx=='60FE':current='对象已注册但结构有差异：VAR:00，标准子项缺失'
    xj,ev=declared(idx)
    result='高创列出，我方缺对象' if xp and not op else ('双方均列出，需比属性/行为' if xp else ('我方有，高创XML未列' if op else '双方资料均未列，不计参考分母'))
    inventory.append([v['section'],f'0x{idx}',v['name'],v['category'],int(op),int(xp),current,xj,ev,result,ours[idx]['ref'] if op else 'df-sdk/modules/ethercat/cia402appl.h:1673–1745；coeappl.c:494–512'])
overlap=sum(r[4] and r[5] for r in inventory);denom=sum(r[5] for r in inventory)
assert (overlap,denom)==(23,72)
issues=[]
for i,r in enumerate(changed['对象属性逐项']['rows'],7):
    if '有差异' in r[7]:issues.append([r[0],r[1],r[3],r[4],r[7],r[8],f'对象属性逐项 第{i}行'])
issues.insert(0,['F07-B04','8.4.12','6502 bit4','客户PDF图7要求保留位为0','0x1DD的bit4=1；同时HM bit5=0但回零对象已列',evidence('6502',0),'行为与位定义逐项'])
source=dict(path=str(xml),sha256=hashlib.sha256(xml.read_bytes()).hexdigest(),device=d.findtext('Type'),product=d.find('Type').get('ProductCode'),revision=d.find('Type').get('RevisionNo'),vendor=doc.findtext('Vendor/Id'),profile=d.findtext('Profile/ProfileNo'),objects=len(objects),mask=mask,coe=dict(d.find('Mailbox/CoE').attrib))
result=dict(changed=changed,inventory=inventory,issues=issues,source=source,objects=objects,stats=dict(standard=len(standard),reference=denom,ours=overlap,missing=denom-overlap,ourOnly=[r[1] for r in inventory if r[4] and not r[5]]))
(cache/'cdhd-comparison.json').write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding='utf-8')
print('Source:',source);print('Coverage:',result['stats']);print('Declaration differences:',len(issues))
for r in issues:print(r[:5])
